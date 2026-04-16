from django.contrib import admin
from django.http import HttpResponse
from django.utils.html import format_html

from .models import ExpenseCategory, Topup, Expense, ExpenseAttachment, AuditLog


@admin.register(ExpenseCategory)
class ExpenseCategoryAdmin(admin.ModelAdmin):
    list_display = ("name", "is_active", "sort_order")
    list_editable = ("is_active", "sort_order")
    search_fields = ("name",)


@admin.register(Topup)
class TopupAdmin(admin.ModelAdmin):
    list_display = ("date", "user", "amount", "source", "comment", "created_by", "created_at")
    list_filter = ("user", "date", "source")
    search_fields = ("comment",)
    date_hierarchy = "date"

    def save_model(self, request, obj, form, change):
        if not change and not obj.created_by_id:
            obj.created_by = request.user
        if not change and not obj.source:
            obj.source = "admin"
        super().save_model(request, obj, form, change)


class ExpenseAttachmentInline(admin.TabularInline):
    model = ExpenseAttachment
    extra = 0
    readonly_fields = ("original_name", "mime_type", "size", "uploaded_at", "preview_link")

    def preview_link(self, obj):
        if obj.pk:
            return format_html('<a href="/attachments/{}/preview/" target="_blank">Просмотр</a>', obj.pk)
        return "—"
    preview_link.short_description = "Файл"


@admin.register(Expense)
class ExpenseAdmin(admin.ModelAdmin):
    list_display = (
        "date", "user", "category", "custom_category_name",
        "amount", "comment_short", "created_via", "is_deleted", "created_at",
    )
    list_filter = ("user", "category", "created_via", "is_deleted", "date")
    search_fields = ("comment", "custom_category_name")
    date_hierarchy = "date"
    inlines = [ExpenseAttachmentInline]
    actions = ["export_as_xlsx"]

    def comment_short(self, obj):
        return (obj.comment or "")[:40]
    comment_short.short_description = "Комментарий"

    def export_as_xlsx(self, request, queryset):
        import openpyxl
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Расходы"

        headers = ["Дата", "Сотрудник", "Категория", "Сумма (₸)", "Комментарий", "Источник", "Вложения"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = openpyxl.styles.Font(bold=True)

        for row_num, e in enumerate(queryset.select_related("user", "category").prefetch_related("attachments"), 2):
            ws.cell(row=row_num, column=1, value=e.date.isoformat())
            ws.cell(row=row_num, column=2, value=e.user.username)
            ws.cell(row=row_num, column=3, value=e.category.name if e.category else e.custom_category_name)
            ws.cell(row=row_num, column=4, value=float(e.amount))
            ws.cell(row=row_num, column=5, value=e.comment or "")
            ws.cell(row=row_num, column=6, value=e.created_via)

            attachments = e.attachments.all()
            if attachments:
                links = []
                for att in attachments:
                    url = f"/attachments/{att.id}/preview/"
                    links.append(f"{att.original_name}: {request.build_absolute_uri(url)}")
                ws.cell(row=row_num, column=7, value="\n".join(links))
            else:
                ws.cell(row=row_num, column=7, value="—")

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="expenses.xlsx"'
        wb.save(response)
        return response
    export_as_xlsx.short_description = "Экспорт в Excel"


@admin.register(AuditLog)
class AuditLogAdmin(admin.ModelAdmin):
    list_display = ("created_at", "user", "action", "entity", "entity_id")
    list_filter = ("action", "entity", "user")
    readonly_fields = ("user", "action", "entity", "entity_id", "diff", "created_at")

    def has_add_permission(self, request):
        return False

    def has_change_permission(self, request, obj=None):
        return False

    def has_delete_permission(self, request, obj=None):
        return False
